"""Export catalogue/moteurs.csv as an .xlsx laid out like the Google Sheet
(one tab per brand, same columns), ready for File > Import in Google Sheets.

Usage: python tools/export_sheet.py [catalogue/moteurs_complete.xlsx]
"""
import csv, re, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "catalogue" / "moteurs.csv"
COLS = ["REF", "MARQUE", "NOM", "VERSION", "CLASSE", "KV", "POIDS", "H STATOR", "D STATOR", "H MOTEUR", "D MOTEUR",
        "D SHAFT", "L SHAFT", "TYPE SHAFT", "VIS HEL", "VIS FIX", "ENTRAXE FIX", "LIPO", "VOLTAGE", "L CABLE",
        "TYPE CABLE", "HELICE", "PUISSANCE", "AMP", "AIMANT", "CLOCHE", "CONFIG", "LIEN", "IMG", "RESISTANCE", "UTILISATION"]
NUMERIC = {"KV", "POIDS", "H MOTEUR", "D MOTEUR", "D SHAFT", "L SHAFT", "PUISSANCE", "AMP"}


def value(col, v):
    if col in NUMERIC and re.fullmatch(r"\d+(\.\d+)?", v or ""):
        return float(v) if "." in v else int(v)
    return v or None


def main(out):
    with SRC.open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    wb = Workbook()
    wb.remove(wb.active)
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="222222")
    for brand in sorted({r["MARQUE"] for r in rows}, key=str.lower):
        ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "-", brand.upper())[:31])
        ws.append(COLS)
        for c in ws[1]:
            c.font, c.fill = head, fill
        for r in sorted((r for r in rows if r["MARQUE"] == brand), key=lambda r: (r.get("CLASSE", ""), r.get("KV", ""))):
            ws.append([value(c, r.get(c, "")) for c in COLS])
        ws.freeze_panes = "B2"
        for i, c in enumerate(COLS, 1):
            ws.column_dimensions[get_column_letter(i)].width = 28 if c in ("LIEN", "IMG", "NOM") else 12
    wb.save(out)
    print(f"{len(rows)} moteurs, {len(wb.sheetnames)} onglets -> {out}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ROOT / "catalogue" / "moteurs_complete.xlsx")
